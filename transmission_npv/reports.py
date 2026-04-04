"""Output formatting and summary reports."""

from __future__ import annotations

from typing import TYPE_CHECKING

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from transmission_npv.model import TransmissionNPVModel


def format_summary(summary: dict) -> str:
    """Format model summary as a readable text report."""
    lines = [
        "=" * 65,
        f"  TRANSMISSION NPV MODEL - {summary['project_name']}",
        "=" * 65,
        "",
        "PROJECT TIMELINE",
        f"  Construction:     {summary['start_year']} - {summary['cod_year'] - 1} ({summary['cod_year'] - summary['start_year']} years)",
        f"  Operations:       {summary['cod_year']} - {summary['end_year']} ({summary['end_year'] - summary['cod_year'] + 1} years)",
        f"  Total project:    {summary['total_years']} years",
        "",
        "CAPITAL EXPENDITURE",
        f"  Initial CAPEX:    ${summary['total_initial_capex_$M']:,.1f}M",
        f"  Lifecycle CAPEX:  ${summary['total_lifecycle_capex_$M']:,.1f}M",
        f"  Peak RAB:         ${summary['peak_rab_$M']:,.1f}M",
        "",
        "COST OF CAPITAL",
        f"  Post-tax WACC:    {summary['wacc_post_tax']:.2%}",
        f"  Pre-tax WACC:     {summary['wacc_pre_tax']:.2%}",
        f"  Regulatory WACC:  {summary['regulatory_wacc']:.2%}",
        "",
        "VALUATION METRICS",
        f"  Project NPV:      ${summary['npv_project_$M']:,.1f}M",
        f"  Equity NPV:       ${summary['npv_equity_$M']:,.1f}M",
    ]

    if summary.get("irr_project") is not None:
        lines.append(f"  Project IRR:      {summary['irr_project']:.2%}")
    if summary.get("irr_equity") is not None:
        lines.append(f"  Equity IRR:       {summary['irr_equity']:.2%}")
    if summary.get("simple_payback_years") is not None:
        lines.append(f"  Simple Payback:   {summary['simple_payback_years']} years")

    lines += [
        "",
        "REVENUE & COSTS (LIFETIME)",
        f"  Total Revenue:    ${summary['total_revenue_$M']:,.1f}M",
        f"  Total OPEX:       ${summary['total_opex_$M']:,.1f}M",
        f"  Total Tax:        ${summary['total_tax_$M']:,.1f}M",
        "",
        "=" * 65,
    ]
    return "\n".join(lines)


def format_cashflow_table(df: pd.DataFrame, years: int = 10) -> str:
    """Format first N years of cash flow as a text table."""
    subset = df.head(years)
    cols = ["revenue", "opex", "capex", "fcff", "fcfe"]
    available = [c for c in cols if c in subset.columns]
    table = subset[available].copy()

    for col in available:
        table[col] = table[col].apply(lambda x: f"${x:,.1f}M")

    return table.to_string()


def export_to_csv(df: pd.DataFrame, path: str) -> None:
    """Export DataFrame to CSV."""
    df.to_csv(path)


def export_to_excel(model: TransmissionNPVModel, path: str) -> None:
    """Export full model to a formatted multi-sheet Excel workbook.

    Sheets:
        1. Summary - Key metrics dashboard
        2. Cash Flows - Year-by-year FCFF/FCFE
        3. CAPEX - Capital expenditure schedule
        4. RAB - Regulated Asset Base roll-forward
        5. OPEX - Operating expenditure breakdown
        6. Revenue - Allowed revenue components
        7. Debt - Debt schedule and interest
        8. Tax - Tax calculation
        9. Parameters - Full model configuration
    """
    from openpyxl import Workbook

    wb = Workbook()

    # --- Styles ---
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    title_font = Font(name="Calibri", bold=True, size=14, color="2F5496")
    section_font = Font(name="Calibri", bold=True, size=11, color="2F5496")
    number_font = Font(name="Calibri", size=11)
    money_fmt = '#,##0.0'
    pct_fmt = '0.00%'
    int_fmt = '#,##0'
    thin_border = Border(
        bottom=Side(style="thin", color="D9D9D9"),
    )
    header_border = Border(
        bottom=Side(style="medium", color="2F5496"),
    )

    def style_header_row(ws, row, max_col):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = header_border

    def write_df_sheet(ws, df, money_columns=None, pct_columns=None):
        """Write a DataFrame to a worksheet with formatting."""
        money_columns = money_columns or []
        pct_columns = pct_columns or []

        # Headers
        ws.cell(row=1, column=1, value="Year")
        for j, col in enumerate(df.columns, start=2):
            ws.cell(row=1, column=j, value=col.replace("_", " ").title())
        style_header_row(ws, 1, len(df.columns) + 1)

        # Data
        for i, (idx, row) in enumerate(df.iterrows(), start=2):
            ws.cell(row=i, column=1, value=idx).font = Font(bold=True, size=11)
            for j, col in enumerate(df.columns, start=2):
                cell = ws.cell(row=i, column=j, value=row[col])
                cell.font = number_font
                cell.border = thin_border
                if col in money_columns:
                    cell.number_format = money_fmt
                elif col in pct_columns:
                    cell.number_format = pct_fmt
                else:
                    cell.number_format = money_fmt

        # Column widths
        ws.column_dimensions["A"].width = 8
        for j in range(2, len(df.columns) + 2):
            ws.column_dimensions[get_column_letter(j)].width = 16

        # Freeze header
        ws.freeze_panes = "B2"

    # =============================================
    # Sheet 1: Summary
    # =============================================
    ws = wb.active
    ws.title = "Summary"
    summary = model.summary()

    ws.cell(row=1, column=1, value="TRANSMISSION NPV MODEL").font = Font(
        name="Calibri", bold=True, size=18, color="2F5496"
    )
    ws.merge_cells("A1:D1")
    ws.cell(row=2, column=1, value=summary["project_name"]).font = Font(
        name="Calibri", size=12, italic=True, color="666666"
    )
    ws.merge_cells("A2:D2")

    sections = [
        ("PROJECT TIMELINE", [
            ("Construction Period", f"{summary['start_year']} - {summary['cod_year'] - 1}"),
            ("Operating Period", f"{summary['cod_year']} - {summary['end_year']}"),
            ("Total Project Life", f"{summary['total_years']} years"),
        ]),
        ("CAPITAL EXPENDITURE ($M)", [
            ("Initial CAPEX", summary["total_initial_capex_$M"]),
            ("Lifecycle CAPEX", summary["total_lifecycle_capex_$M"]),
            ("Peak RAB", summary["peak_rab_$M"]),
        ]),
        ("COST OF CAPITAL", [
            ("Post-tax WACC", summary["wacc_post_tax"]),
            ("Pre-tax WACC", summary["wacc_pre_tax"]),
            ("Regulatory WACC", summary["regulatory_wacc"]),
        ]),
        ("VALUATION METRICS", [
            ("Project NPV ($M)", summary["npv_project_$M"]),
            ("Equity NPV ($M)", summary["npv_equity_$M"]),
            ("Project IRR", summary.get("irr_project")),
            ("Equity IRR", summary.get("irr_equity")),
            ("Simple Payback (years)", summary.get("simple_payback_years")),
        ]),
        ("LIFETIME TOTALS ($M)", [
            ("Total Revenue", summary["total_revenue_$M"]),
            ("Total OPEX", summary["total_opex_$M"]),
            ("Total Tax", summary["total_tax_$M"]),
        ]),
    ]

    row = 4
    for section_title, items in sections:
        ws.cell(row=row, column=1, value=section_title).font = section_font
        ws.merge_cells(f"A{row}:B{row}")
        row += 1
        for label, value in items:
            ws.cell(row=row, column=1, value=label).font = number_font
            cell = ws.cell(row=row, column=2, value=value)
            cell.font = Font(name="Calibri", bold=True, size=11)
            if isinstance(value, float):
                if "WACC" in section_title or "IRR" in label:
                    cell.number_format = pct_fmt
                else:
                    cell.number_format = money_fmt
            row += 1
        row += 1  # blank row between sections

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18

    # =============================================
    # Sheet 2: Cash Flows
    # =============================================
    ws_cf = wb.create_sheet("Cash Flows")
    cf = model.cashflows
    write_df_sheet(ws_cf, cf)

    # =============================================
    # Sheet 3: CAPEX
    # =============================================
    ws_capex = wb.create_sheet("CAPEX")
    write_df_sheet(ws_capex, model.capex_schedule)

    # =============================================
    # Sheet 4: RAB
    # =============================================
    ws_rab = wb.create_sheet("RAB")
    write_df_sheet(ws_rab, model.rab_schedule)

    # =============================================
    # Sheet 5: OPEX
    # =============================================
    ws_opex = wb.create_sheet("OPEX")
    write_df_sheet(ws_opex, model.opex_projection)

    # =============================================
    # Sheet 6: Revenue
    # =============================================
    ws_rev = wb.create_sheet("Revenue")
    write_df_sheet(ws_rev, model.revenue)

    # =============================================
    # Sheet 7: Debt
    # =============================================
    ws_debt = wb.create_sheet("Debt")
    write_df_sheet(ws_debt, model.debt_schedule)

    # =============================================
    # Sheet 8: Tax
    # =============================================
    ws_tax = wb.create_sheet("Tax")
    write_df_sheet(ws_tax, model.tax_schedule)

    # =============================================
    # Sheet 9: Parameters
    # =============================================
    ws_params = wb.create_sheet("Parameters")
    ws_params.cell(row=1, column=1, value="Parameter").font = header_font
    ws_params.cell(row=1, column=1).fill = header_fill
    ws_params.cell(row=1, column=2, value="Value").font = header_font
    ws_params.cell(row=1, column=2).fill = header_fill
    style_header_row(ws_params, 1, 2)

    config_dict = model.config.to_dict()
    row = 2
    for section_name, section_data in config_dict.items():
        ws_params.cell(row=row, column=1, value=section_name.upper()).font = section_font
        row += 1
        if isinstance(section_data, dict):
            for key, value in section_data.items():
                ws_params.cell(row=row, column=1, value=f"  {key}").font = number_font
                cell = ws_params.cell(row=row, column=2, value=str(value))
                cell.font = number_font
                row += 1
        row += 1

    ws_params.column_dimensions["A"].width = 35
    ws_params.column_dimensions["B"].width = 25

    # Save
    wb.save(path)


def format_scenario_comparison(comparison_df: pd.DataFrame) -> str:
    """Format scenario comparison table."""
    display_cols = [
        "npv_project", "irr_project", "total_capex", "total_revenue"
    ]
    available = [c for c in display_cols if c in comparison_df.columns]
    table = comparison_df[available].copy()

    for col in available:
        if "npv" in col or "capex" in col or "revenue" in col:
            table[col] = table[col].apply(lambda x: f"${x:,.1f}M")
        elif "irr" in col:
            table[col] = table[col].apply(
                lambda x: f"{x:.2%}" if x is not None else "N/A"
            )

    lines = [
        "=" * 65,
        "  SCENARIO COMPARISON",
        "=" * 65,
        "",
        table.to_string(),
        "",
        "=" * 65,
    ]
    return "\n".join(lines)
